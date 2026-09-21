"""Finans raporunu Excel (.xlsx) olarak dışa aktarır. Sekmeye göre (scope) yalnızca ilgili sayfaları üretir ve
ekrandaki filtreleri (dönem, ülke, arama, sıralama) dosyaya yansıtır; her sayfanın en üstünde filtre notu bulunur."""
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.finance import service
from app.shops.models import Shop

HEAD_FILL = PatternFill("solid", fgColor="F1641E")
MONEY = "#,##0.00"
SORT_LABELS = {"sales": "satışa göre", "profit": "kâra göre", "margin": "marja göre"}


def _sheet(wb: Workbook, title: str, note: str, headers: list[str], rows: list[list], money_cols: set[int], widths: dict[int, int] | None = None) -> None:
    ws = wb.create_sheet(title)
    ws.append([note])
    ws["A1"].font = Font(italic=True, color="666666")
    ws.append(headers)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for idx in money_cols:
        for cell in ws[get_column_letter(idx)][2:]:
            cell.number_format = MONEY
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(12, min(40, len(h) + 4)))
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{max(2, ws.max_row)}"


def build_xlsx(
    db: Session,
    shop: Shop,
    start: dt.date,
    end: dt.date,
    country: str = "",
    scope: str = "all",
    q: str = "",
    sort: str = "sales",
) -> bytes:
    """scope: all | overview | products | orders"""
    order_rows: list[dict] = []
    rep = service.report(db, shop, start, end, country, collect=order_rows)
    cur = rep["currency"]
    k = rep["kpi"]
    want = {"all": {"overview", "products", "orders"}}.get(scope, {scope})

    def note(*extra: str) -> str:
        parts = [f"Dönem: {start} – {end}", f"Ülke: {country or 'Tümü'}", f"Para birimi: {cur}", *extra]
        return "Filtre — " + " · ".join(parts)

    wb = Workbook()
    wb.remove(wb.active)

    if "overview" in want:
        ws = wb.create_sheet("Özet")
        ws.append([note()])
        ws["A1"].font = Font(italic=True, color="666666")
        ws.append([])
        ws.append(["Toplam", "Tutar"])
        for c in ws[3]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
        for label, v in [
            ("Sipariş sayısı", k["orders"]),
            ("Satış (vergi hariç)", k["sales"]),
            ("İadeler", k["refunds"]),
            ("Sipariş ücretleri", k["fees"]),
            *[(f"  {n}", v) for n, v in k["fee_types"].items()],
            ("Reklam / yenileme / diğer giderler", k["overhead"]),
            *[(f"  {n}", v) for n, v in k["overhead_types"].items()],
            ("Ürün + kargo maliyeti", k["cogs"]),
            ("Net kâr", k["profit"]),
            ("Kâr marjı (%)", round(k["margin"], 2)),
        ]:
            ws.append([label, v])
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 18
        for row in ws.iter_rows(min_row=4, min_col=2, max_col=2):
            row[0].number_format = MONEY
        ws.append([])
        base = ws.max_row + 1
        ws.append(["Ay", "Satış", "Sipariş ücretleri", "Reklam/diğer", "Ürün+kargo", "Net kâr", "Sipariş", "Geçen yıl satış", "Geçen yıl net kâr"])
        for c in ws[base]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
        for sp in rep["series"]:
            ws.append([sp["month"], sp["sales"], sp["fees"], sp["overhead"], sp["cogs"], sp["profit"], sp["orders"], sp["prev_sales"], sp["prev_profit"]])
        for row in ws.iter_rows(min_row=base + 1, min_col=2, max_col=9):
            for i, cell in enumerate(row):
                if i != 5:
                    cell.number_format = MONEY
        for col in "CDEFGHI":
            ws.column_dimensions[col].width = 18

        _sheet(
            wb, "Müşteriler", note(), ["Müşteri", "Ülke", "Sipariş", "Toplam satış", "Son sipariş"],
            [[c["name"], c["country"], c["orders"], c["sales"], c["last"]] for c in rep["customers"]], {4}, {1: 30},
        )
        _sheet(
            wb, "Ülkeler", note(), ["Ülke", "Sipariş", "Satış", "Geçen yıl sipariş", "Geçen yıl satış"],
            [[c["iso"], c["orders"], c["sales"], c["prev_orders"], c["prev_sales"]] for c in rep["countries"]], {3, 5},
        )

    if "products" in want:
        products = sorted(rep["products"], key=lambda p: -p.get(sort if sort in SORT_LABELS else "sales", 0))
        _sheet(
            wb, "Ürünler", note(f"Sıralama: {SORT_LABELS.get(sort, 'satışa göre')}"),
            ["Listing ID", "Ürün", "Seçenek", "Adet", "Satış", "Etsy ücreti", "İade", "Maliyet (ürün+kargo)", "Kalan", "Marj %"],
            [
                row
                for p in products
                for row in (
                    [[p["listing_id"], p["title"], "(tümü)", p["units"], p["sales"], p["fees"], p["refunds"], p["cogs"], p["profit"], round(p["margin"], 1)]]
                    + [[p["listing_id"], p["title"], v["key"] or "Seçeneksiz", v["units"], v["sales"], v.get("fees"), None, v["cogs"], v["sales"] - v.get("fees", 0) - v["cogs"], None] for v in p["variants"]]
                )
            ],
            {5, 6, 7, 8, 9}, {2: 50, 3: 45},
        )

    if "orders" in want:
        needle = q.strip().lower()
        rows = [
            o for o in order_rows
            if not needle or needle in f"{o['buyer']} {o['items']} {o['receipt_id']} {o['country']}".lower()
        ]
        _sheet(
            wb, "Siparişler", note(f"Arama: {q.strip()}" if needle else "Arama: yok", f"{len(rows)} sipariş"),
            [
                "Sipariş no", "Tarih", "Müşteri", "Ülke", "Ürünler", "Sipariş toplamı", "Vergi", "Satış (vergi hariç)", "İade", "İşlem ücreti",
                "Ödeme işleme ücreti", "Düzenleyici ücret", "Diğer ücret", "Maliyet", "Maliyet elle girildi", "Kalan (kâr)", "Ücret kaydı var",
            ],
            [
                [
                    o["receipt_id"], o["date"], o["buyer"], o["country"], o["items"], o["grand"], o["tax"], o["sales"], o["refunds"], o["transaction_fee"],
                    o["processing_fee"], o["regulatory_fee"], o["other_fee"], o["cogs"], "Evet" if o["manual_cost"] else "", o["profit"], "Evet" if o["fees_known"] else "Hayır",
                ]
                for o in sorted(rows, key=lambda r: r["date"])
            ],
            {6, 7, 8, 9, 10, 11, 12, 13, 14, 16}, {3: 26, 5: 60},
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
