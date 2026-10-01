"""Kargo faturalarından (PDF/görsel/Excel/CSV/HTML) veri çıkarımı ve SİPARİŞLE eşleştirme.

Bir kargo faturası ürün adı taşımaz; her satırı bir GÖNDERİdir (takip no, alıcı adı, ülke, ağırlık, tutar).
Bu yüzden eşleştirme siparişe yapılır: önce takip numarası (kesin), olmazsa alıcı adı + ülke + tarih. Tutar,
eşleşen siparişin ürünlerine fiyat payına göre dağıtılıp `ShippingInvoice` satırlarına yazılır ve o siparişin
kargo maliyeti olarak (elle girilen tahminin yerine) kullanılır.

Akış: `parse_file()` -> "aday" satırlar (henüz KAYDEDİLMEZ, önizleme) -> kullanıcı onaylar -> `confirm()`.
Dosyanın kendisi saklanmaz; yalnızca çıkarılan sayısal veri.
"""
import csv
import datetime as dt
import html
import io
import json
import logging
import re
from collections import defaultdict

import httpx
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.finance.models import ShippingInvoice
from app.orders.models import OrderCache
from app.shops.models import Shop

logger = logging.getLogger("app.finance.invoices")

KINDS = ("nakliye", "gümrük", "ek hizmet", "diğer")


class InvoiceError(Exception):
    pass


# --------------------------------------------------------------- çıkarım ----

EXTRACT_PROMPT = (
    "Bu bir kargo/nakliye faturası. Faturadan yalnızca şu JSON'u çıkar, başka hiçbir metin ekleme:\n"
    '{"vendor": "firma (ör. FedEx)", "invoice_number": "fatura no", "invoice_date": "YYYY-MM-DD", '
    '"currency": "fatura para birimi kodu (TL için TRY)", '
    '"total_amount_text": "faturanın genel/ödenecek toplam tutarı, faturada yazdığı gibi metin", '
    '"stated_rate": faturada bir döviz kuru yazıyorsa {"base": "USD", "quote": "TRY", "rate": 48.8741} '
    "(rate: 1 base = kaç quote; ör. 'USD:48.8741TL' -> base USD, quote TRY, rate 48.8741), yoksa null, "
    '"shipments": [{"tracking_no": "gönderi/takip numarası", "ship_date": "YYYY-MM-DD", '
    '"recipient": "ALICI adı-soyadı (gönderen değil)", "recipient_country": "alıcı ülkesi", '
    '"reference": "referans no", "weight_kg": sayı ya da null, '
    '"amount_text": "tutar TAM OLARAK faturada yazdığı gibi metin (ör. \'4.007,19 TL\', \'1,250.00\', \'1 250\')", "amount": sayı, '
    '"kind": "nakliye" | "gümrük" | "ek hizmet" | "diğer", "description": "kalemin faturadaki adı (ör. Hizmet Ücreti, Gümrük Vergisi)"}]}\n'
    "KURALLAR: (1) shipments'e her GÖNDERİ SATIRI için ayrı bir kayıt yaz. (2) Fatura toplamlarını, KDV/vergi "
    "satırlarını ve 'dahil edilmiştir' denen ücretleri (ör. EPH ücreti tutara dahilse) AYRI kalem olarak EKLEME — "
    "çift sayılır. (3) amount_text ve total_amount_text'e tutarı faturada yazdığı gibi AYNEN kopyala (binlik/ondalık ayraçlarını değiştirme, yorumlama); weight_kg düz sayı olsun (ör. '8,5 KG' -> 8.5). (4) Emin olmadığını null bırak, "
    "uydurma. (5) Bir gönderi için birden fazla ücret kalemi varsa (ör. nakliye + gümrük vergisi + ek hizmet ücreti) HER "
    "KALEMİ ayrı kayıt olarak yaz, aynı takip no ile. Tür: nakliye = taşıma bedeli; gümrük = gümrük vergisi/işlem; ek hizmet = "
    "standart dışı paket/ekstra hizmet ücreti; diğer = geri kalanı."
)


def _extract_json(text: str) -> dict:
    match = re.search(r"\{.*\}", text, re.DOTALL)
    try:
        return json.loads(match.group(0) if match else text)
    except (json.JSONDecodeError, AttributeError) as exc:
        raise InvoiceError("Yapay zekâ faturadan okunabilir bir sonuç döndürmedi, tekrar dene.") from exc


def extract_pdf_or_image(content: bytes, content_type: str) -> dict:
    """Taranmış PDF ya da fotoğraf — vision/dosya destekli AI ile okunur (pahalı yol; yalnızca metin katmanı yoksa
    ya da metin yolu güvenilir sonuç vermezse). Claude ve OpenAI ikisi de PDF'i doğrudan okuyabilir."""
    import base64

    from app.ai.client import get_anthropic_client, get_openai_client
    from app.ai.vision import VisionError, _provider
    from app.core.config import settings

    try:
        provider = _provider()
    except VisionError as exc:
        raise InvoiceError(str(exc)) from exc

    b64 = base64.b64encode(content).decode()
    try:
        if provider == "anthropic":
            media_type = content_type if content_type in ("image/png", "image/jpeg", "image/webp", "image/gif") else "application/pdf"
            doc_type = "image" if media_type.startswith("image/") else "document"
            resp = get_anthropic_client().messages.create(
                model=settings.anthropic_model,
                max_tokens=2000,
                messages=[{"role": "user", "content": [
                    {"type": doc_type, "source": {"type": "base64", "media_type": media_type, "data": b64}},
                    {"type": "text", "text": EXTRACT_PROMPT},
                ]}],
            )
            raw = "".join(b.text for b in resp.content if b.type == "text")
        else:
            if content_type == "application/pdf":  # OpenAI PDF'i "dosya" parçası olarak okur (resim değil)
                part = {"type": "file", "file": {"filename": "fatura.pdf", "file_data": f"data:application/pdf;base64,{b64}"}}
            else:
                part = {"type": "image_url", "image_url": {"url": f"data:{content_type};base64,{b64}"}}
            resp = get_openai_client().chat.completions.create(
                model=settings.openai_model,
                messages=[{"role": "user", "content": [{"type": "text", "text": EXTRACT_PROMPT}, part]}],
                max_tokens=2000,
            )
            raw = resp.choices[0].message.content or ""
    except Exception as exc:  # noqa: BLE001
        raise InvoiceError(f"Yapay zekâ faturayı okuyamadı: {str(exc)[:200]}") from exc
    return _extract_json(raw)


def _extract_from_text(text: str) -> dict:
    """Metin (PDF metin katmanı / HTML) -> aynı çıkarım promptu. Görüntü değil yalnızca metin gittiği için ucuzdur."""
    from app.ai.client import get_anthropic_client, get_openai_client
    from app.ai.vision import VisionError, _provider
    from app.core.config import settings

    try:
        provider = _provider()
    except VisionError as exc:
        raise InvoiceError(str(exc)) from exc
    prompt = EXTRACT_PROMPT + "\n\nFatura metni:\n" + text[:12000]
    try:
        if provider == "anthropic":
            resp = get_anthropic_client().messages.create(model=settings.anthropic_model, max_tokens=2000, messages=[{"role": "user", "content": prompt}])
            raw = "".join(b.text for b in resp.content if b.type == "text")
        else:
            resp = get_openai_client().chat.completions.create(model=settings.openai_model, messages=[{"role": "user", "content": prompt}], max_tokens=2000)
            raw = resp.choices[0].message.content or ""
    except Exception as exc:  # noqa: BLE001
        raise InvoiceError(f"Yapay zekâ faturayı okuyamadı: {str(exc)[:200]}") from exc
    return _extract_json(raw)


def extract_html(content: bytes) -> dict:
    text = re.sub(r"<[^>]+>", " ", content.decode("utf-8", errors="ignore"))
    return _extract_from_text(html.unescape(re.sub(r"\s+", " ", text)).strip())


MIN_PDF_TEXT_CHARS = 80  # bunun altı = metin katmanı yok (taranmış PDF); görüntü yoluna düşülür


def pdf_text(content: bytes) -> str:
    """PDF'in metin katmanını yerelde (ücretsiz, AI'sız) çıkarır; şifreli/bozuk/taranmış PDF'te boş döner."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return re.sub(r"\s+", " ", " ".join((page.extract_text() or "") for page in reader.pages)).strip()
    except Exception:  # noqa: BLE001
        return ""


MAX_TEXT_CHARS = 12000  # bundan uzun metin AI'ya kesilerek gitmez (satır kaybı riski); PDF doğrudan gönderilir


def _usable(raw: dict) -> bool:
    """Metin yolunun sonucu güvenilir mi: en az bir gönderi satırı ve hepsinde sayısal tutar var."""
    lines = raw.get("shipments") or []
    return bool(lines) and all(_to_float(x.get("amount")) for x in lines)


def extract_pdf(content: bytes) -> dict:
    """Önce PDF'in TÜM metni (yerel, ucuz) AI'ya verilir — faturalar firmaya/düzene göre değiştiği için kural
    tabanlı ayrıştırma yerine AI okur. Risk varsa (metin yok/çok uzun, AI hata verdi ya da sonuç eksik) PDF
    doğrudan gönderilir: daha pahalı ama en güvenilir yol."""
    text = pdf_text(content)
    if MIN_PDF_TEXT_CHARS <= len(text) <= MAX_TEXT_CHARS:
        try:
            raw = _extract_from_text(text)
            if _usable(raw):
                return raw
            logger.info("PDF metin yolu eksik sonuç verdi; PDF doğrudan gönderiliyor")
        except InvoiceError:
            logger.info("PDF metin yolu başarısız; PDF doğrudan gönderiliyor")
    return extract_pdf_or_image(content, "application/pdf")


# Excel/CSV: AI'sız — sütun başlıklarından doğrudan okunur.
COLUMN_ALIASES = {
    "tracking_no": ("tracking", "takip", "gönderi no", "awb", "waybill"),
    "recipient": ("recipient", "alıcı", "alici", "consignee", "name", "isim"),
    "recipient_country": ("country", "ülke", "ulke", "destination"),
    "ship_date": ("ship date", "date", "tarih"),
    "amount": ("amount", "total", "tutar", "ücret", "charge", "fiyat", "price"),
    "currency": ("currency", "para birimi", "ccy"),
    "kind": ("kind", "type", "tür", "kategori"),
    "weight_kg": ("weight", "ağırlık", "agirlik", "kg"),
    "description": ("description", "servis", "service", "açıklama"),
}


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9çğıöşü]+", " ", str(h).lower()).strip()


def _match_column(headers: list[str], field: str) -> int | None:
    normed = [_norm_header(h) for h in headers]
    for i, h in enumerate(normed):
        if any(a in h for a in COLUMN_ALIASES[field]):
            return i
    return None


def extract_spreadsheet(content: bytes, filename: str) -> dict:
    """Excel (xlsx/xls) ya da CSV — her satır bir gönderi olur."""
    lower = filename.lower()
    if lower.endswith(".csv"):
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig", errors="ignore"))))
    elif lower.endswith(".xls"):
        try:
            import xlrd
        except ImportError as exc:
            raise InvoiceError("Eski .xls formatı için xlrd paketi kurulu değil.") from exc
        sheet = xlrd.open_workbook(file_contents=content).sheet_by_index(0)
        rows = [sheet.row_values(i) for i in range(sheet.nrows)]
    else:
        from openpyxl import load_workbook

        rows = [list(r) for r in load_workbook(io.BytesIO(content), read_only=True, data_only=True).active.iter_rows(values_only=True)]

    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise InvoiceError("Dosyada okunabilir satır bulunamadı.")
    headers = [str(c or "") for c in rows[0]]
    idx = {f: _match_column(headers, f) for f in COLUMN_ALIASES}
    if idx["amount"] is None:
        raise InvoiceError("Dosyada bir 'tutar' sütunu bulunamadı (Amount/Total/Tutar gibi bir başlık bekleniyor).")

    def cell(row: list, key: str):
        i = idx[key]
        return row[i] if i is not None and i < len(row) else None

    shipments, currency = [], None
    for row in rows[1:]:
        amount = _to_float(cell(row, "amount"))
        if not amount:
            continue
        currency = currency or (str(cell(row, "currency") or "").upper() or None)
        shipments.append({
            "tracking_no": str(cell(row, "tracking_no") or "").strip().split(".")[0],
            "ship_date": _parse_date_cell(cell(row, "ship_date")),
            "recipient": str(cell(row, "recipient") or ""),
            "recipient_country": str(cell(row, "recipient_country") or ""),
            "reference": "",
            "weight_kg": _to_float(cell(row, "weight_kg")),
            "amount": amount,
            "kind": _norm_kind(cell(row, "kind")),
            "description": str(cell(row, "description") or ""),
        })
    if not shipments:
        raise InvoiceError("Dosyada geçerli tutarlı hiçbir satır bulunamadı.")
    return {"vendor": "", "invoice_number": "", "invoice_date": None, "currency": currency, "stated_rate": None, "shipments": shipments}


def parse_amount(v) -> tuple[float | None, bool]:
    """Para tutarı okuyucu: 1250 · 1,250 · 1.250 · 1 250 · 1'250 · 1.250,50 · 1,250.50 · ₺/$/TL/USD gibi eklerle.
    Döner: (değer, belirsiz mi). İki ayraç varsa sonuncusu ondalıktır; tek ayraçta ve ardında TAM 3 hane varsa
    ("1,250" / "1.250") para tutarında 3 ondalık olmayacağından binlik sayılır ama BELİRSİZ işaretlenir (fatura toplamıyla
    çapraz kontrol edilir); 1-2 hane ondalıktır; aynı ayraç birden fazlaysa binliktir."""
    import re as _re

    if v is None or v == "":
        return None, False
    if isinstance(v, (int, float)):
        return float(v), False
    s = _re.sub(r"[^\d,.\-\s'\u2019\u00a0]", "", str(v))
    neg = "-" in s
    s = _re.sub(r"[\s'\u2019\u00a0-]", "", s)
    if not s or not any(ch.isdigit() for ch in s):
        return None, False
    ambiguous = False
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif "," in s or "." in s:
        sep = "," if "," in s else "."
        parts = s.split(sep)
        if len(parts) > 2:
            s = "".join(parts)
        else:
            head, tail = parts
            if len(tail) == 3 and head not in ("", "0") and len(head) <= 3:
                s, ambiguous = head + tail, True
            else:
                s = (head or "0") + "." + tail
    try:
        val = float(s)
    except ValueError:
        return None, False
    return (-val if neg else val), ambiguous


def _to_float(v) -> float | None:
    return parse_amount(v)[0]


def _line_amounts(raw: dict) -> tuple[list[float | None], str]:
    """Her gönderi satırının tutarı (yazıldığı metinden okunur) + fatura toplamıyla çapraz kontrol. Belirsiz ayraçlı
    tutarlar ("1,250") için satır toplamı fatura toplamına daha yakın olan yorum seçilir. Toplam tutmuyorsa uyarı notu döner."""
    parsed = []
    for ln in raw.get("shipments") or []:
        text = ln.get("amount_text")
        val, amb = parse_amount(text if text not in (None, "") else ln.get("amount"))
        if val is None:
            val, amb = parse_amount(ln.get("amount"))
        parsed.append([val, amb])
    ttext = raw.get("total_amount_text")
    total, total_amb = parse_amount(ttext if ttext not in (None, "") else raw.get("total_amount"))
    valid = [p for p in parsed if p[0]]
    note = ""
    if total and not total_amb and valid:
        base = sum(p[0] for p in valid)
        if any(p[1] for p in valid):
            alt = sum(p[0] / 1000 if p[1] else p[0] for p in valid)
            if abs(alt - total) < abs(base - total):
                for p in parsed:
                    if p[1] and p[0]:
                        p[0] /= 1000
                base = alt
        if abs(base - total) > max(0.02 * total, 0.05):
            note = f"Satırların toplamı ({base:.2f}) faturanın toplamından ({total:.2f}) farklı; tutarları kontrol edin."
    return [p[0] for p in parsed], note


def _norm_kind(v) -> str:
    s = str(v or "").strip().lower()
    if "gümrük" in s or "customs" in s or "duty" in s:
        return "gümrük"
    if "hizmet" in s or "surcharge" in s or "service fee" in s or "extra" in s or "ek " in s:
        return "ek hizmet"
    if "nakliye" in s or "freight" in s or "shipping" in s:
        return "nakliye"
    return "diğer"


def _parse_date_cell(v) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(str(v).strip(), fmt).date().isoformat()
        except ValueError:
            continue
    return None


# --------------------------------------------------------------- sipariş eşleştirme ----

_WORD = re.compile(r"[a-z0-9]+")
_TR = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")

COUNTRY_ISO = {
    "united states": "US", "usa": "US", "canada": "CA", "united kingdom": "GB", "uk": "GB", "germany": "DE",
    "france": "FR", "australia": "AU", "italy": "IT", "spain": "ES", "netherlands": "NL", "turkey": "TR",
    "türkiye": "TR", "ireland": "IE", "sweden": "SE", "norway": "NO", "denmark": "DK", "switzerland": "CH",
    "belgium": "BE", "austria": "AT", "new zealand": "NZ", "mexico": "MX", "japan": "JP",
}


def _name_words(text: str) -> set[str]:
    return {w for w in _WORD.findall((text or "").translate(_TR).lower()) if len(w) > 1}


def _digits(s: str) -> str:
    return re.sub(r"\D", "", s or "")


def _order_index(db: Session, shop: Shop) -> list[dict]:
    """Eşleştirme için siparişlerin hafif özeti: takip kodları, alıcı, ülke, tarih ve kalemler."""
    out = []
    for row in db.scalars(select(OrderCache).where(OrderCache.shop_id == shop.id)):
        raw = json.loads(row.raw_json)
        out.append({
            "receipt_id": row.receipt_id,
            "canceled": bool(row.is_canceled),
            "buyer": row.buyer_name or "",
            "country": (row.country_iso or "").upper(),
            "date": row.created_at.date(),
            "tracking": {_digits(s.get("tracking_code")) for s in raw.get("shipments") or [] if s.get("tracking_code")},
            "txs": raw.get("transactions") or [],
        })
    return out


def _order_items(o: dict) -> list[dict]:
    from app.finance.service import variant_key

    return [
        {"listing_id": t.get("listing_id"), "variant_key": variant_key(t), "title": html.unescape(t.get("title") or "")[:80],
         "qty": t.get("quantity") or 1}
        for t in o["txs"]
    ]


def match_orders(index: list[dict], line: dict, limit: int = 3) -> list[dict]:
    """Bir fatura satırını siparişlerle eşleştirir. Takip no birebir tutarsa kesin (1.0); değilse alıcı adı +
    ülke + tarih yakınlığından bir güven skoru üretilir."""
    tn = _digits(line.get("tracking_no"))
    ship_date = None
    if line.get("ship_date"):
        try:
            ship_date = dt.date.fromisoformat(line["ship_date"])
        except ValueError:
            pass
    iso = COUNTRY_ISO.get((line.get("recipient_country") or "").strip().lower())
    rw = _name_words(line.get("recipient"))
    scored = []
    for o in index:
        if tn and tn in o["tracking"]:
            scored.append((1.0, "takip no", o))
            continue
        bw = _name_words(o["buyer"])
        if not rw or not bw:
            continue
        name = len(rw & bw) / max(len(rw), len(bw))
        if name < 0.5:
            continue
        country = 1.0 if iso and o["country"] == iso else (0.0 if iso else 0.5)
        near = 0.5
        if ship_date:
            gap = (ship_date - o["date"]).days
            near = 1.0 if 0 <= gap <= 30 else 0.0
        # Ad benzerliği "kesin" sayılamaz (aynı isimde farklı müşteri olabilir): tavan 0.89, yalnızca takip no 1.0 verir.
        scored.append((round(min(0.89, 0.6 * name + 0.25 * country + 0.15 * near), 2), "alıcı adı + ülke", o))
    scored.sort(key=lambda x: -x[0])
    if scored and scored[0][0] >= 1.0:  # takip no birebir tuttu: benzer isimli diğer siparişler yalnızca kafa karıştırır
        scored = [x for x in scored if x[0] >= 1.0]
    return [
        {"receipt_id": o["receipt_id"], "canceled": o["canceled"], "buyer": o["buyer"], "country": o["country"], "date": o["date"].isoformat(),
         "score": s, "reason": why, "items": _order_items(o)}
        for s, why, o in scored[:limit]
    ]


# --------------------------------------------------------------- kur çevrimi ----


def resolve_fx(currency: str | None, amount: float, on: str | None, report_ccy: str, stated: dict | None) -> tuple[float, float, str]:
    """Döner: (rapor para biriminde tutar, "1 fatura birimi = kaç rapor birimi", kaynak).
    Öncelik: faturada yazan kur (yönü doğru uygulanır) > aynı para birimi > tarihe göre geçmiş kur (Frankfurter)."""
    currency = (currency or report_ccy).upper()
    report_ccy = report_ccy.upper()
    if currency == report_ccy:
        return amount, 1.0, "aynı para birimi"
    if stated and stated.get("rate"):
        try:
            base, quote, rate = str(stated["base"]).upper(), str(stated["quote"]).upper(), float(stated["rate"])
            if rate > 0 and currency == quote and report_ccy == base:  # ör. TRY faturayı USD'ye: bölünür
                return amount / rate, 1 / rate, "faturada yazan kur"
            if rate > 0 and currency == base and report_ccy == quote:
                return amount * rate, rate, "faturada yazan kur"
        except (KeyError, TypeError, ValueError):
            pass
    date_str = on or dt.date.today().isoformat()
    try:
        resp = httpx.get(f"https://api.frankfurter.dev/v1/{date_str}", params={"base": currency, "symbols": report_ccy}, timeout=10, follow_redirects=True)
        resp.raise_for_status()
        rate = resp.json()["rates"][report_ccy]
        return amount * rate, rate, f"{date_str} tarihli kur"
    except Exception:  # noqa: BLE001
        logger.warning("Fatura kuru alınamadı (%s -> %s, %s)", currency, report_ccy, date_str)
        return amount, 1.0, "kur alınamadı — 1.0 kabul edildi, kontrol et"


# --------------------------------------------------------------- üst seviye ----


def _stem(name: str) -> str:
    return (name or "").rsplit(".", 1)[0]


def fingerprint(invoice_no: str, tracking: str, receipt_id: int | None, kind: str, original_amount: float) -> str:
    """Bir fatura kaleminin kimliği: fatura no + gönderi + tür + tutar. Aynı faturayı tekrar yüklemek engellenir; aynı
    gönderiye kesilen FARKLI kalemler (farklı tutar/tür/fatura) ayrı kayıt olarak girebilir."""
    who = tracking or f"r{receipt_id}"
    return f"{invoice_no}|{who}|{kind}|{original_amount:.2f}"


def parse_file(db: Session, shop: Shop, content: bytes, filename: str, content_type: str, report_ccy: str) -> list[dict]:
    """Bir dosyayı okuyup, KAYDETMEDEN, her gönderi satırı için bir "aday" döner (onay ekranı için)."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xls", ".csv")):
        raw = extract_spreadsheet(content, filename)
    elif lower.endswith((".html", ".htm")):
        raw = extract_html(content)
    elif content_type == "application/pdf" or lower.endswith(".pdf"):
        raw = extract_pdf(content)
    elif content_type.startswith("image/") or lower.endswith((".jpg", ".jpeg", ".png", ".webp")):
        raw = extract_pdf_or_image(content, content_type)
    else:
        raise InvoiceError(f"Desteklenmeyen dosya türü: {filename}")

    index = _order_index(db, shop)
    saved_fp = set(db.scalars(select(ShippingInvoice.fingerprint).where(ShippingInvoice.shop_id == shop.id)))
    invoice_date = raw.get("invoice_date") or dt.date.today().isoformat()
    candidates = []
    amounts, check_note = _line_amounts(raw)
    for line, amount in zip(raw.get("shipments") or [], amounts):
        if not amount:
            continue
        converted, fx_rate, fx_source = resolve_fx(raw.get("currency"), amount, invoice_date, report_ccy, raw.get("stated_rate"))
        kind = line.get("kind") if line.get("kind") in KINDS else "nakliye"
        tracking = _digits(line.get("tracking_no"))
        invoice_no = (raw.get("invoice_number") or "").strip() or _stem(filename)
        candidates.append({
            "vendor": raw.get("vendor") or "",
            "invoice_number": invoice_no,
            "invoice_date": invoice_date,
            "kind": kind,
            "description": line.get("description") or "",
            "tracking_no": tracking,
            "ship_date": line.get("ship_date"),
            "recipient": line.get("recipient") or "",
            "recipient_country": line.get("recipient_country") or "",
            "weight_kg": _to_float(line.get("weight_kg")),
            "original_amount": amount,
            "check_note": check_note,
            "original_currency": (raw.get("currency") or report_ccy).upper(),
            "amount": round(converted, 2),
            "fx_rate": fx_rate,
            "fx_source": fx_source,
            "matches": match_orders(index, line),
            "already_saved": bool(tracking) and fingerprint(invoice_no, tracking, None, kind, amount) in saved_fp,
            "source_filename": filename,
        })
    if not candidates:
        raise InvoiceError("Faturadan hiçbir gönderi satırı çıkarılamadı.")
    return candidates


def confirm(db: Session, shop: Shop, receipt_id: int, candidate: dict) -> list[ShippingInvoice]:
    """Onaylanan gönderi kalemini, siparişin ürünlerine fiyat payına göre dağıtıp kaydeder."""
    from app.finance.service import _money, variant_key

    row = db.scalar(select(OrderCache).where(OrderCache.shop_id == shop.id, OrderCache.receipt_id == receipt_id))
    if row is None:
        raise InvoiceError("Sipariş bulunamadı.")
    tracking = _digits(candidate.get("tracking_no"))
    kind = candidate.get("kind") if candidate.get("kind") in KINDS else "nakliye"
    invoice_no = (str(candidate.get("invoice_number") or "").strip() or _stem(str(candidate.get("source_filename") or "")))[:60]
    original = float(candidate["original_amount"])
    fp = fingerprint(invoice_no, tracking, receipt_id, kind, original)
    if db.scalar(select(ShippingInvoice.id).where(ShippingInvoice.shop_id == shop.id, ShippingInvoice.fingerprint == fp)):
        raise InvoiceError("Bu fatura kalemi zaten kayıtlı (aynı fatura no, gönderi, tür ve tutar).")
    txs = json.loads(row.raw_json).get("transactions") or []
    if not txs:
        raise InvoiceError("Siparişte ürün kalemi yok.")
    physical = [t for t in txs if not t.get("is_digital")]
    if not physical:
        raise InvoiceError("Dijital ürünlerden oluşan siparişe kargo faturası yazılamaz.")
    txs = physical  # dijital kalemin kargosu olmaz; tutar yalnızca fiziksel kalemlere dağıtılır (aksi halde payı sessizce kaybolurdu)
    weights = [_money(t.get("price")) * (t.get("quantity") or 1) for t in txs]
    total_w = sum(weights) or float(len(txs))
    amount = float(candidate["amount"])
    out = []
    for t, w in zip(txs, weights):
        share = (w or 1.0) / total_w
        item = ShippingInvoice(
            shop_id=shop.id, receipt_id=receipt_id, tracking_no=tracking,
            listing_id=t.get("listing_id") or 0, variant_key=variant_key(t)[:300], kind=kind,
            description=str(candidate.get("description") or "")[:200], invoice_no=invoice_no, fingerprint=fp,
            amount=round(amount * share, 4), original_amount=round(original * share, 4),
            original_currency=str(candidate.get("original_currency") or "")[:10],
            fx_rate=float(candidate.get("fx_rate") or 1.0), fx_source=str(candidate.get("fx_source") or "")[:20],
            invoice_date=dt.date.fromisoformat(candidate["invoice_date"]),
            weight_kg=round(float(candidate["weight_kg"]) * share, 3) if candidate.get("weight_kg") else None,
            vendor=str(candidate.get("vendor") or "")[:120],
            source_filename=str(candidate.get("source_filename") or "")[:255],
            match_confidence=float(next((m["score"] for m in candidate.get("matches") or [] if m["receipt_id"] == receipt_id), 0.0)),
        )
        db.add(item)
        out.append(item)
    db.commit()
    return out


def list_invoices(db: Session, shop: Shop, listing_id: int | None = None) -> list[ShippingInvoice]:
    stmt = select(ShippingInvoice).where(ShippingInvoice.shop_id == shop.id)
    if listing_id is not None:
        stmt = stmt.where(ShippingInvoice.listing_id == listing_id)
    return list(db.scalars(stmt.order_by(ShippingInvoice.invoice_date.desc(), ShippingInvoice.id)))


def _lines(rows: list[ShippingInvoice]) -> list[dict]:
    """Aynı parmak izli (aynı kalemin ürünlere dağıtılmış) satırları tek kaleme birleştirir."""
    by_fp: dict[str, list[ShippingInvoice]] = {}
    for r in rows:
        by_fp.setdefault(r.fingerprint or f"id{r.id}", []).append(r)
    lines = []
    for items in by_fp.values():
        f = items[0]
        lines.append({
            "id": f.id, "kind": f.kind, "description": f.description, "invoice_no": f.invoice_no,
            "invoice_date": f.invoice_date.isoformat(), "amount": round(sum(i.amount for i in items), 2),
            "original_amount": round(sum(i.original_amount for i in items), 2), "original_currency": f.original_currency,
            "fx_source": f.fx_source, "source_filename": f.source_filename,
            "weight_kg": round(sum(i.weight_kg or 0 for i in items), 3) or None,
        })
    lines.sort(key=lambda x: (x["invoice_date"], x["id"]))
    return lines


def _warnings(lines: list[dict]) -> list[str]:
    """Aynı gönderiye şüpheli/fazla kalem kesilmiş mi: kullanıcının hatalı faturayı yakalaması için."""
    out = []
    if len(lines) >= 3:
        out.append(f"Bu gönderiye {len(lines)} kalem kesilmiş")
    kinds: dict[str, int] = {}
    for ln in lines:
        kinds[ln["kind"]] = kinds.get(ln["kind"], 0) + 1
    out += [f"{n} adet '{k}' kalemi" for k, n in kinds.items() if n >= 2]
    names: dict[str, int] = {}
    for ln in lines:
        key = (ln["description"] or "").strip().lower()
        if key:
            names[key] = names.get(key, 0) + 1
    out += [f"'{k}' kalemi {n} kez kesilmiş" for k, n in names.items() if n >= 2]
    return out


def query_invoices(
    db: Session, shop: Shop, q: str = "", kind: str = "", inv_start: str = "", inv_end: str = "",
    order_start: str = "", order_end: str = "", sort: str = "inv_date", page: int = 0, per_page: int = 20,
) -> dict:
    """Kayıtlı faturaları GÖNDERİ başına tek kayıt olarak döner (içinde kalemler: nakliye/gümrük/ek hizmet…). Arama:
    müşteri, takip no, ürün, kalem adı, fatura no. Süzgeçler: tür, fatura/sipariş tarihi. Küçük tablo, bellekte süzülür."""
    from app.listings.models import ListingCache

    rows = list_invoices(db, shop)
    rids = {r.receipt_id for r in rows if r.receipt_id}
    order_info: dict[int, tuple[str, dt.date]] = {}
    if rids:
        for rid, name, created in db.execute(select(OrderCache.receipt_id, OrderCache.buyer_name, OrderCache.created_at).where(OrderCache.shop_id == shop.id, OrderCache.receipt_id.in_(rids))):
            order_info[rid] = (name or "", created.date())
    titles = {lid: t for lid, t in db.execute(select(ListingCache.listing_id, ListingCache.title).where(ListingCache.shop_id == shop.id))}

    shipments: dict[tuple, list[ShippingInvoice]] = {}
    for r in rows:
        shipments.setdefault((r.receipt_id, r.tracking_no or f"id{r.id}"), []).append(r)

    needle = q.strip().lower()
    d = lambda v: dt.date.fromisoformat(v) if v else None  # noqa: E731
    i0, i1, o0, o1 = d(inv_start), d(inv_end), d(order_start), d(order_end)
    items = []
    for (rid, trk), grp in shipments.items():
        buyer, odate = order_info.get(rid, ("", None))
        lines = _lines(grp)
        if kind and not any(ln["kind"] == kind for ln in lines):
            continue
        dates = [dt.date.fromisoformat(ln["invoice_date"]) for ln in lines]
        if i0 and max(dates) < i0 or i1 and min(dates) > i1:
            continue
        if (o0 or o1) and (odate is None or (o0 and odate < o0) or (o1 and odate > o1)):
            continue
        products = []
        for r in grp:
            key = (r.listing_id, r.variant_key)
            if key not in [(p["listing_id"], p["variant_key"]) for p in products]:
                products.append({"listing_id": r.listing_id, "variant_key": r.variant_key, "title": titles.get(r.listing_id, "")})
        blob = " ".join([buyer, trk, *(p["title"] + " " + p["variant_key"] for p in products), *(f"{ln['description']} {ln['invoice_no']} {ln['kind']}" for ln in lines), grp[0].vendor]).lower()
        if needle and needle not in blob:
            continue
        items.append({
            "receipt_id": rid, "tracking_no": grp[0].tracking_no, "buyer": buyer, "order_date": odate.isoformat() if odate else None,
            "vendor": grp[0].vendor, "products": products, "lines": lines,
            "total": round(sum(ln["amount"] for ln in lines), 2), "weight_kg": max((ln["weight_kg"] or 0 for ln in lines), default=0) or None,
            "last_invoice_date": max(dates).isoformat(), "warnings": _warnings(lines),
        })
    keys = {
        "inv_date": (lambda x: x["last_invoice_date"], True),
        "order_date": (lambda x: x["order_date"] or "", True),
        "buyer": (lambda x: x["buyer"].lower(), False),
        "amount": (lambda x: x["total"], True),
    }
    fn, rev = keys.get(sort, keys["inv_date"])
    items.sort(key=fn, reverse=rev)
    total = len(items)
    per_page = max(1, min(per_page, 200))
    return {
        "items": items[page * per_page : (page + 1) * per_page],
        "total": total,
        "total_amount": round(sum(x["total"] for x in items), 2),
    }


def order_lines(db: Session, shop: Shop, receipt_id: int) -> list[dict]:
    """Bir siparişe yazılmış tüm kargo/gümrük/ek hizmet kalemleri (Sipariş maliyetleri dökümü için)."""
    rows = list(db.scalars(select(ShippingInvoice).where(ShippingInvoice.shop_id == shop.id, ShippingInvoice.receipt_id == receipt_id)))
    return _lines(rows)


def buyer_names(db: Session, shop: Shop, receipt_ids: set[int]) -> dict[int, str]:
    """Kayıtlı fatura satırlarında müşteri adını göstermek için: sipariş no -> alıcı adı (ayrıca saklanmaz)."""
    if not receipt_ids:
        return {}
    rows = db.execute(select(OrderCache.receipt_id, OrderCache.buyer_name).where(OrderCache.shop_id == shop.id, OrderCache.receipt_id.in_(receipt_ids)))
    return {rid: name or "" for rid, name in rows}


def delete_invoice(db: Session, shop: Shop, invoice_id: int) -> None:
    """Bir kalem birden fazla ürüne dağıtılmış olabilir; aynı parmak izli (aynı kalemin) tüm satırlar birlikte silinir."""
    row = db.get(ShippingInvoice, invoice_id)
    if row is None or row.shop_id != shop.id:
        return
    stmt = select(ShippingInvoice).where(ShippingInvoice.shop_id == shop.id)
    stmt = stmt.where(ShippingInvoice.fingerprint == row.fingerprint) if row.fingerprint else stmt.where(ShippingInvoice.id == row.id)
    for r in db.scalars(stmt).all():
        db.delete(r)
    db.commit()


def delete_invoices(db: Session, shop: Shop, invoice_ids: list[int]) -> int:
    """Toplu silme: verilen kalemlerin (her biri için dağıtılmış tüm satırlarıyla) hepsi tek işlemde silinir. Döner: silinen kalem sayısı."""
    fingerprints, single_ids = set(), set()
    for iid in invoice_ids:
        row = db.get(ShippingInvoice, iid)
        if row is None or row.shop_id != shop.id:
            continue
        (fingerprints if row.fingerprint else single_ids).add(row.fingerprint or row.id)
    stmt = select(ShippingInvoice).where(ShippingInvoice.shop_id == shop.id)
    rows = []
    if fingerprints:
        rows += db.scalars(stmt.where(ShippingInvoice.fingerprint.in_(fingerprints))).all()
    if single_ids:
        rows += db.scalars(stmt.where(ShippingInvoice.id.in_(single_ids))).all()
    for r in rows:
        db.delete(r)
    db.commit()
    return len(fingerprints) + len(single_ids)


def order_shipping(db: Session, shop: Shop) -> dict[tuple[int, int, str], float]:
    """(receipt_id, listing_id, variant_key) -> o siparişin o kalemine yazılan GERÇEK kargo tutarı (rapor para biriminde).
    `service.report` bunu, elle girilmiş tahmini kargo maliyetinin yerine kullanır."""
    out: dict[tuple[int, int, str], float] = defaultdict(float)
    for r in db.scalars(select(ShippingInvoice).where(ShippingInvoice.shop_id == shop.id, ShippingInvoice.receipt_id.is_not(None))):
        out[(r.receipt_id, r.listing_id, r.variant_key)] += r.amount
    return dict(out)


def invoice_totals(db: Session, shop: Shop) -> dict[tuple[int, str], dict]:
    """(listing_id, variant_key) -> {"amount": toplam, "weight_kg": ortalama ağırlık, "count": gönderi adedi} — yalnızca
    Ürün kârlılığı tablosunda bilgi göstermek için."""
    out: dict[tuple[int, str], dict] = defaultdict(lambda: {"amount": 0.0, "count": 0, "_w": [], "_t": set()})
    for row in db.scalars(select(ShippingInvoice).where(ShippingInvoice.shop_id == shop.id)):
        agg = out[(row.listing_id, row.variant_key)]
        agg["amount"] += row.amount
        agg["_t"].add((row.tracking_no or f"id{row.id}", row.kind))
        if row.weight_kg:
            agg["_w"].append(row.weight_kg)
    for agg in out.values():
        agg["count"] = len(agg.pop("_t"))
        w = agg.pop("_w")
        agg["weight_kg"] = round(sum(w) / len(w), 3) if w else None
    return dict(out)
